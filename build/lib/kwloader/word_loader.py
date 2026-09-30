from io import BytesIO
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import BookDocument, TextBlock, TextRun, TextStyle


_WORD_DOCUMENT_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)
_WORD_TEMPLATE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)


def _points(value):
    if value is None:
        return None
    try:
        return float(value.pt)
    except Exception:
        return None


def _style_font(paragraph):
    try:
        font = paragraph.style.font
        return TextStyle(
            bold=bool(font.bold),
            italic=bool(font.italic),
            underline=bool(font.underline),
            size=_points(font.size),
            font=font.name,
        )
    except Exception:
        return TextStyle()


def _load_dotx(path: Path):
    """
    python-docx identifies the main part by its OOXML content type.
    A .dotx uses the template content type, so normalize that one XML
    content-type entry in memory before handing the package to python-docx.
    The original file is never modified.
    """
    buffer = BytesIO()

    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        buffer, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)

            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    _WORD_TEMPLATE_CONTENT_TYPE.encode(),
                    _WORD_DOCUMENT_CONTENT_TYPE.encode(),
                )

            target.writestr(item, data)

    buffer.seek(0)
    return Document(buffer)


def _load_document(path: Path):
    if path.suffix.lower() == ".dotx":
        return _load_dotx(path)

    return Document(str(path))


def _paragraph_alignment(paragraph):
    value = paragraph.alignment

    if value == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    if value == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    if value == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return "justify"
    if value == WD_ALIGN_PARAGRAPH.LEFT:
        return "left"

    return None


def _paragraph_kind(paragraph):
    try:
        style_name = paragraph.style.name or ""
    except Exception:
        style_name = ""

    lowered = style_name.lower()
    if lowered.startswith("heading") or lowered in {
        "title",
        "subtitle",
    }:
        return "heading"

    return "paragraph"


def load_word(path: str | Path) -> BookDocument:
    path = Path(path).resolve()
    document = _load_document(path)

    blocks: list[TextBlock] = []

    for paragraph in document.paragraphs:
        if not paragraph.text.strip() and not paragraph.runs:
            continue

        base_style = _style_font(paragraph)
        runs: list[TextRun] = []

        for run in paragraph.runs:
            font = run.font
            style = TextStyle(
                bold=bool(run.bold) if run.bold is not None else base_style.bold,
                italic=(
                    bool(run.italic)
                    if run.italic is not None
                    else base_style.italic
                ),
                underline=(
                    bool(run.underline)
                    if run.underline is not None
                    else base_style.underline
                ),
                size=(
                    _points(font.size)
                    if font.size is not None
                    else base_style.size
                ),
                font=font.name or base_style.font,
            )

            if run.text:
                runs.append(TextRun(run.text, style))

        # Some Word paragraphs can have text without a useful run list.
        if not runs and paragraph.text:
            runs.append(TextRun(paragraph.text, base_style))

        if not runs:
            continue

        blocks.append(
            TextBlock(
                kind=_paragraph_kind(paragraph),
                runs=runs,
                align=_paragraph_alignment(paragraph),
                style_name=getattr(paragraph.style, "name", None),
            )
        )

    title = path.stem.replace("_", " ").replace("-", " ").title()

    for block in blocks:
        if block.kind == "heading" and block.text.strip():
            title = block.text.strip()
            break

    if title == path.stem.replace("_", " ").replace("-", " ").title():
        for block in blocks:
            if block.text.strip():
                title = block.text.strip()
                break

    return BookDocument(
        path=path,
        title=title,
        format=path.suffix.lower(),
        blocks=blocks,
    )
